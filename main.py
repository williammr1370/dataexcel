import io
import os
import tempfile
from typing import List

import openpyxl
from openpyxl.styles import PatternFill
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware

app = FastAPI()

# Permitir solicitudes desde el frontend React
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # En producción, pon la URL de tu React
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

verde_claro = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

def convertir_a_segundos(valor):
    if valor is None:
        return 0
    if hasattr(valor, "hour"):
        return valor.hour * 3600 + valor.minute * 60 + valor.second
    if isinstance(valor, str):
        partes = valor.strip().split(":")
        try:
            if len(partes) == 3:
                return int(partes[0]) * 3600 + int(partes[1]) * 60 + int(partes[2])
            elif len(partes) == 2:
                return int(partes[0]) * 60 + int(partes[1])
        except ValueError:
            return 0
    if isinstance(valor, (int, float)):
        if valor < 1:
            return int(round(valor * 86400))
        return int(valor)
    return 0

def es_vacio(valor):
    return valor is None or (isinstance(valor, str) and valor.strip() == "")

@app.post("/api/procesar-excel")
async def procesar_excels(files: List[UploadFile] = File(...)):
    if not files:
        raise HTTPException(status_code=400, detail="No se subieron archivos.")

    periodo_detectado = None
    archivos_validos = []

    # Validar nombres de archivos subidos
    for file in files:
        nombre_file = file.filename
        nombre_base, _ = os.path.splitext(nombre_file)

        if len(nombre_base) > 33 and nombre_base[32] == "_":
            periodo = nombre_base[33:]
            archivos_validos.append(file)
            if periodo_detectado is None:
                periodo_detectado = periodo

    if not archivos_validos:
        raise HTTPException(
            status_code=400,
            detail="Ningún archivo cumple con el formato [keyname(32ch)]_[periodo].xlsx"
        )

    # 1. Unificar
    wb_unificado = openpyxl.Workbook()
    ws_unificado = wb_unificado.active
    ws_unificado.title = "Unificado"

    encabezados_guardados = False
    row_destino = 1

    for file in archivos_validos:
        contenido = await file.read()
        wb_temp = openpyxl.load_workbook(filename=io.BytesIO(contenido), data_only=False)
        ws_temp = wb_temp.active

        max_c = ws_temp.max_column
        max_r = ws_temp.max_row

        start_r = 1
        if not encabezados_guardados:
            for c in range(1, max_c + 1):
                ws_unificado.cell(row=1, column=c, value=ws_temp.cell(row=1, column=c).value)
            encabezados_guardados = True
            start_r = 2
            row_destino = 2
        else:
            start_r = 2

        for r in range(start_r, max_r + 1):
            if any(ws_temp.cell(row=r, column=c).value is not None for c in range(1, max_c + 1)):
                for c in range(1, max_c + 1):
                    celda_src = ws_temp.cell(row=r, column=c)
                    celda_dst = ws_unificado.cell(row=row_destino, column=c, value=celda_src.value)
                    if celda_src.hyperlink:
                        celda_dst.hyperlink = celda_src.hyperlink.target
                row_destino += 1

        wb_temp.close()

    # 2. Formatear y evaluar condicionales
    headers = {
        str(ws_unificado.cell(row=1, column=c).value).strip(): c
        for c in range(1, ws_unificado.max_column + 1)
    }

    cols_requeridas = ["CBlock", "M/S", "Duration", "BmatId"]
    for col_req in cols_requeridas:
        if col_req not in headers:
            raise HTTPException(
                status_code=400,
                detail=f"Falta la columna requerida '{col_req}' en los archivos."
            )

    col_cblock = headers["CBlock"]
    col_ms = headers["M/S"]
    col_duration = headers["Duration"]
    col_bmatid = headers["BmatId"]

    for row in range(2, ws_unificado.max_row + 1):
        val_ms = ws_unificado.cell(row=row, column=col_ms).value
        val_duration = ws_unificado.cell(row=row, column=col_duration).value
        val_bmatid = ws_unificado.cell(row=row, column=col_bmatid).value

        es_music = str(val_ms).strip().lower() == "music" if val_ms is not None else False
        duracion_ok = convertir_a_segundos(val_duration) >= 120
        bmatid_vacio = es_vacio(val_bmatid)

        if es_music and duracion_ok and bmatid_vacio:
            ws_unificado.cell(row=row, column=col_cblock).fill = verde_claro

    # 3. Guardar en un archivo temporal para retornarlo
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb_unificado.save(temp_file.name)
    wb_unificado.close()

    nombre_salida = f"CRP_{periodo_detectado}_procesado.xlsx"
    return FileResponse(
        path=temp_file.name,
        filename=nombre_salida,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )